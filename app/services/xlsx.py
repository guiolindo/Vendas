"""Exportação para Excel (.xlsx de verdade): números, datas e percentuais viram tipos do Excel
(somam, filtram e ordenam), não texto. Um relatório = uma planilha; o "pacote" junta várias."""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import clock
from ..domain.status import LABELS, Status
from .reports import Report

INK = "1C2622"
PAPER = "F3F0E8"
LINE = "D9D2C3"
MONEY_FMT = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
FORMATS = {"money": MONEY_FMT, "int": "#,##0", "percent": "0.0%", "date": "dd/mm/yyyy", "link": '"#"0'}
INVALID_TITLE = set('[]:*?/\\')


def _sheet_title(text: str, used: set[str]) -> str:
    base = "".join("-" if c in INVALID_TITLE else c for c in text)[:31] or "Planilha"
    title, n = base, 2
    while title.lower() in used:
        title = f"{base[:28]} {n}"; n += 1
    used.add(title.lower())
    return title


def _value(kind: str, v):
    if v is None or v == "":
        return None
    if kind == "money":
        return Decimal(int(v)) / 100
    if kind == "percent":
        return Decimal(str(v)) / 100
    if kind in ("int", "link"):
        return int(v)
    if kind == "status":
        return LABELS[Status(v)]
    return v


def write_sheet(ws: Worksheet, report: Report, subtitle: str = "") -> None:
    cols = report.columns
    ws.sheet_view.showGridLines = False
    ws["A1"] = report.title
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=INK)
    ws["A2"] = subtitle or report.description
    ws["A2"].font = Font(name="Calibri", size=10, color="7B8680")
    header_row = 4
    thin = Side(style="thin", color=LINE)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=header_row, column=j, value=c.label)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="right" if c.kind in ("money", "int", "percent") else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 22
    r = header_row
    for i, row in enumerate(report.rows):
        r += 1
        for j, c in enumerate(cols, 1):
            raw = row.get(c.key)
            cell = ws.cell(row=r, column=j)
            cell.value = _value(c.kind, raw)
            if isinstance(cell.value, str):
                cell.data_type = "s"   # texto é texto: nunca vira fórmula (=, +, -, @ no começo)
            if c.kind in FORMATS and cell.value is not None:
                cell.number_format = FORMATS[c.kind]
            cell.alignment = Alignment(horizontal="right" if c.kind in ("money", "int", "percent") else "left", vertical="center")
            cell.border = Border(bottom=thin)
            if i % 2:
                cell.fill = PatternFill("solid", fgColor="FAF8F2")
    if report.totals:
        r += 1
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=r, column=j)
            if j == 1:
                cell.value = "Total"
            elif c.key in report.totals:
                cell.value = _value(c.kind, report.totals[c.key])
                cell.number_format = FORMATS.get(c.kind, "General")
            cell.font = Font(name="Calibri", bold=True, color=INK)
            cell.fill = PatternFill("solid", fgColor=PAPER)
            cell.border = Border(top=Side(style="medium", color=INK))
            cell.alignment = Alignment(horizontal="right" if c.kind in ("money", "int", "percent") else "left")
    # largura pelo conteúdo (com teto), cabeçalho congelado, filtro, impressão em paisagem
    for j, c in enumerate(cols, 1):
        longest = max([len(str(c.label))] + [len(f"{row.get(c.key)}") for row in report.rows[:200] if row.get(c.key) is not None] + [10])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 12), 46)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if report.rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{header_row + len(report.rows)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"


def to_xlsx(reports: list[Report], subtitle: str = "") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for rep in reports:
        ws = wb.create_sheet(_sheet_title(rep.title, used))
        write_sheet(ws, rep, subtitle)
    wb.properties.creator = "Vendas"
    wb.properties.created = clock.now()
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def period_text(start: date | None, end: date | None) -> str:
    if start and end:
        return f"Período: {start.strftime('%d/%m/%Y')} a {end.strftime('%d/%m/%Y')}"
    return f"Gerado em {clock.today().strftime('%d/%m/%Y')}"
