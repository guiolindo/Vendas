import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.services import reports, xlsx
from app.services import owners as own
from app.services import products as pvc
from app.services import sales as sale_svc
from app.services.sales import ItemInput, SaleInput

from .conftest import TODAY, owned, create_sale_for
from .test_web import app, client, db, post_json  # noqa: F401  (fixtures)


@pytest.fixture
def data(session, make_customer):
    ana = own.create_owner(session, "Ana")
    pa = owned(pvc.create_product(session, pvc.ProductInput(name='=HYPERLINK("http://evil","x")', price_cents=12345, cost_cents=5000, initial_stock=20)), ana)
    pb = owned(pvc.create_product(session, pvc.ProductInput(name="Brinco", price_cents=5000, cost_cents=0, initial_stock=20)), ana)
    c = make_customer("+cmd|' /C calc'!A0")
    create_sale_for(session, SaleInput(items=[ItemInput(pa.id, 2), ItemInput(pb.id, 1)], customer_id=c.id, paid_cents=10000,
                                            payment_method="pix", due_date=TODAY + timedelta(days=5)))
    return pa, pb, c


def load(blob):
    return load_workbook(io.BytesIO(blob))


def test_values_are_real_excel_types_not_text(session, data):
    rep = reports.build(session, "vendas", reports.ReportParams(), TODAY)
    ws = load(xlsx.to_xlsx([rep])).active
    header = [c.value for c in ws[4]]
    assert header[:4] == ["Venda", "Data", "Quem vendeu", "Cliente"]
    row = {h: ws.cell(row=5, column=i + 1) for i, h in enumerate(header)}
    assert row["Venda"].value == 1 and row["Venda"].number_format == '"#"0'
    assert isinstance(row["Data"].value, (date,)) and row["Data"].number_format == "dd/mm/yyyy"
    assert float(row["Total"].value) == 296.90 and "R$" in row["Total"].number_format        # número de verdade: soma e ordena
    assert row["Situação"].value == "Parcial"
    last = ws.cell(row=6, column=1).value
    assert last == "Total" and float(ws.cell(row=6, column=header.index("Total") + 1).value) == 296.90


def test_text_that_looks_like_a_formula_stays_text(session, data):
    rep = reports.build(session, "vendas", reports.ReportParams(), TODAY)
    ws = load(xlsx.to_xlsx([rep])).active
    cell = [c for c in ws[5] if c.value and "cmd" in str(c.value)][0]
    assert cell.data_type == "s" and cell.value == "+cmd|' /C calc'!A0"                      # texto, não fórmula
    rep = reports.build(session, "estoque", reports.ReportParams(), TODAY)
    ws = load(xlsx.to_xlsx([rep])).active
    names = [ws.cell(row=r, column=1) for r in range(5, 7)]
    evil = [c for c in names if c.value and "HYPERLINK" in c.value][0]
    assert evil.data_type == "s" and evil.value.startswith("=HYPERLINK")


def test_percent_money_and_missing_cost_cells(session, data):
    rep = reports.build(session, "margem-produtos", reports.ReportParams(), TODAY)
    ws = load(xlsx.to_xlsx([rep])).active
    header = [c.value for c in ws[4]]
    rows = {ws.cell(row=r, column=1).value: r for r in range(5, 9) if ws.cell(row=r, column=1).value}
    r = rows["Brinco"]
    assert ws.cell(row=r, column=header.index("Custo") + 1).value is None                    # sem custo: célula vazia, não zero
    r = [r for name, r in rows.items() if "HYPERLINK" in name][0]
    pct = ws.cell(row=r, column=header.index("Margem %") + 1)
    assert pct.number_format == "0.0%" and float(pct.value) == pytest.approx(0.5951, abs=1e-3)


def test_sheet_layout_header_frozen_filter_and_print(session, data):
    ws = load(xlsx.to_xlsx([reports.build(session, "vendas", reports.ReportParams(), TODAY)])).active
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref.startswith("A4:")
    assert ws.page_setup.orientation == "landscape" and ws.sheet_view.showGridLines is False
    assert ws["A1"].value == "Vendas por período" and ws["A1"].font.bold


def test_package_has_one_sheet_per_report_with_valid_names(session, data):
    keys = ("pessoas", "vendas", "recebimentos", "a-receber", "produtos", "margem-produtos", "compras", "estoque", "clientes")
    wb = load(xlsx.to_xlsx([reports.build(session, k, reports.ReportParams(date_from=TODAY.replace(day=1), date_to=TODAY), TODAY) for k in keys]))
    assert len(wb.sheetnames) == 9 and len(set(n.lower() for n in wb.sheetnames)) == 9
    assert all(len(n) <= 31 and not set('[]:*?/\\') & set(n) for n in wb.sheetnames)
    assert "Resultado por pessoa" in wb.sheetnames


@pytest.fixture
def web_data(db):
    """Mesmos dados, mas pela sessão do cliente web (no PostgreSQL as duas fixtures compartilhariam o banco)."""
    ana = own.create_owner(db, "Ana")
    pa = owned(pvc.create_product(db, pvc.ProductInput(name="Vestido", price_cents=12345, cost_cents=5000, initial_stock=20)), ana)
    c = __import__("app.services.customers", fromlist=["x"])
    cust = c.create_customer(db, c.CustomerInput(name="Cliente"))
    create_sale_for(db, SaleInput(items=[ItemInput(pa.id, 2)], customer_id=cust.id, paid_cents=10000, payment_method="pix",
                                       due_date=TODAY + timedelta(days=5)))


def test_xlsx_routes(client, db, web_data):
    r = client.get("/relatorios/vendas?exportar=xlsx")
    assert r.status_code == 200 and r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r.data[:2] == b"PK" and ".xlsx" in r.headers["Content-Disposition"]                # um .xlsx é um zip
    assert load(r.data).active["A1"].value == "Vendas por período"
    p = client.get("/relatorios/pacote?de=2026-10-01&ate=2026-10-31")
    assert p.status_code == 200 and len(load(p.data).sheetnames) == 9
    assert client.get("/relatorios/pacote?de=lixo").status_code == 400                       # data inválida: mensagem, nunca 500
    assert "Exportar para Excel" in client.get("/relatorios/vendas").get_data(as_text=True)
    assert "Baixar Excel com tudo" in client.get("/relatorios").get_data(as_text=True)
    for key in reports.BUILDERS:
        assert client.get(f"/relatorios/{key}?exportar=xlsx").status_code == 200, key
